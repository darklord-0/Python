from tkinter import *
from fpdf import FPDF
from docx import Document

root = Tk()

# Save data from Text widget to a text file:
## Conversion to txt file

def save_data():
    text_content = text.get(1.0 , 'end')

    with open("text_file.txt", "w") as f:
        f.write(text_content)
  
def clear():
    text.delete(1.0 , 'end')

sb = Scrollbar(root , )
sb.pack(side=RIGHT,fill=Y , )

text = Text(root ,background='#ccffff' )
text.config(yscrollcommand=sb.set)
text.pack()

sb.config(command=text.yview)

b1 = Button(root, text='1.  GET A TEXT FILE', command=save_data)
b1.pack(padx=5,pady=5)


## Converion to pdf 

def convert():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial' , size=15)

    f =  open('text_file.txt' , 'r' ) 
    for x in f:
        pdf.cell(200 , 10, txt= x  ,ln=1, )
    pdf.output('text_file.pdf') 
    f.close()


convert_file = Button(root , text='2.  CONVERT TO PDF' , command=convert)
convert_file.pack(padx=5,pady=5)


## Conversion to word

def convert2():
    document = Document()
    with open('text_file.txt') as f:
        for line in f:
            document.add_paragraph(line)
    document.save('word_file.docx')

convert_file2 = Button(root , text='3.  CONVERT TO WORD' , command=convert2)
convert_file2.pack(padx=5,pady=5)


b2 = Button(root, text='CLEAR', command=clear)
b2.pack(padx=10,pady=5)



root.mainloop()